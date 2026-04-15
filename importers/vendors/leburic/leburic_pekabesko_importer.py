# importers/leburic_pekabesko_importer.py
"""
ASYCUDA Pro - Leburic/Pekabesko Excel Importer

Parser za Leburic Komerc / Pekabesko AD fakture.

Format Excel fajla (header u redu 1):
  Invoice | Item | Description | Unit | CustomID | Packets | Neto(KGR) | Quantity In Unit | Price Per Unit | Total in EUR

Napomene:
- Zemlja porijekla se čita iz PDF headera (Poreklo: MK) ili default MK
- Bruto težina se čita iz PDF headera (Btol / Btoi) ako postoji PDF u istom folderu
- JM može biti Кгр (Cyrillic, kilogram) ili Пар (Cyrillic, komad/paar)
"""

from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Optional

import openpyxl

from core.draft.draft import InvoiceLine, Party
from importers.import_result import ImportResult
from utils.country_normalizer import normalize_country_name

logger = logging.getLogger("asycuda_pro.import.leburic_pekabesko")

# Normalizacija jedinice mjere (Cyrillic → ASCII)
_JM_MAP = {
    "кгр": "kg",
    "кг": "kg",
    "пар": "par",
    "ком": "kom",
    "кос": "kos",
    "л": "L",
    "мл": "ml",
}

# Regex za ekstrakciju bruto/neto iz PDF teksta (OCR kvalitet je slab)
_BRUTO_RE = re.compile(
    r"bto[il]:?\s*([\d\s\.,]+)\s*k[gs]", re.IGNORECASE
)
_NETO_RE = re.compile(
    r"nto[il]:?\s*([\d\s\.,]+)\s*k[gs]", re.IGNORECASE
)
_POREKLO_RE = re.compile(
    r"poreklo\s*:?\s*([A-Z]{2})", re.IGNORECASE
)


def detect_leburic_pekabesko_excel(filepath: str) -> bool:
    """
    Detekcija Leburic/Pekabesko Excel formata.

    Kriteriji (header red 1):
    - Kolona A = 'Invoice'
    - Kolona B = 'Item'
    - Kolona E = 'CustomID'  (tarifni broj)
    - Kolona G = 'Neto(KGR)'
    """
    try:
        if not filepath.lower().endswith((".xlsx", ".xls")):
            return False

        wb = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
        ws = wb.worksheets[0]

        header = [str(ws.cell(1, c).value or "").strip().lower() for c in range(1, 11)]
        wb.close()

        required = {"invoice", "item", "customid", "neto(kgr)"}
        found = {h for h in header if h in required}

        if len(found) >= 3:
            logger.info(f"Leburic/Pekabesko format detektovan: {filepath}")
            return True
        return False

    except Exception as e:
        logger.debug(f"Greška pri detekciji Leburic/Pekabesko formata: {e}")
        return False


def _read_tariff_code(value) -> str:
    """
    Čita tarifni broj iz Excel ćelije uz čuvanje vodećih nula.

    Excel numeričke ćelije gube vodeće nule: 0210198100 → 210198100.
    TARIC kodovi su uvijek 8 ili 10 cifara, pa:
      9 cifara → paduj na 10 (izgubljena 1 vodeća nula)
      7 cifara → paduj na 8  (izgubljena 1 vodeća nula)
    """
    if value is None:
        return ""
    if isinstance(value, float):
        value = int(value)
    s = str(value).strip().replace(" ", "")
    if not s or s == "0":
        return ""
    if s.isdigit():
        n = len(s)
        if n == 9:
            s = s.zfill(10)
        elif n == 7:
            s = s.zfill(8)
    return s


def _normalize_jm(raw: str) -> str:
    """Normalizuj jedinicu mjere iz Cyrillic u ASCII."""
    s = (raw or "").strip()
    lower = s.lower()
    return _JM_MAP.get(lower, s)


def _parse_number(value) -> float:
    """Parsiraj broj iz Excel ćelije (može biti float ili string sa zarezima)."""
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace(" ", "")
    # Ukloni sve osim cifara, zareza i tačke
    s = re.sub(r"[^\d.,]", "", s)
    if not s:
        return 0.0
    # Ako ima i zarez i tačka, pretpostavi da je zarez separator hiljada
    if "," in s and "." in s:
        s = s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def _find_pdf_for_excel(excel_path: str) -> Optional[str]:
    """
    Pronađi PDF koji odgovara Excel fajlu u istom folderu.

    Matchuje po broju fakture koji se nalazi u imenu Excel fajla.
    Npr. '2000-00015.xlsx' → tražimo PDF čiji tekst sadrži '2000-00015'.
    """
    try:
        import pdfplumber
        folder = Path(excel_path).parent
        stem = Path(excel_path).stem  # npr. '2000-00015'

        for pdf_path in folder.glob("*.pdf"):
            try:
                with pdfplumber.open(str(pdf_path)) as pdf:
                    text = ""
                    for page in pdf.pages[:2]:
                        text += (page.extract_text() or "")
                    if stem in text:
                        logger.info(f"  📄 Nađen matching PDF: {pdf_path.name}")
                        return str(pdf_path)
            except Exception:
                continue
    except Exception as e:
        logger.debug(f"Greška pri traženju PDF para: {e}")
    return None


def _extract_from_pdf(pdf_path: str) -> tuple[float, float, str]:
    """
    Pokušaj izvući bruto, neto i zemlju porijekla iz PDF headera.

    OCR kvalitet je slab, ali Btol/Nto/Poreklo linije su uglavnom čitljive.

    Returns:
        (bruto_kg, neto_kg, zemlja_iso)
    """
    try:
        import pdfplumber
        with pdfplumber.open(pdf_path) as pdf:
            text = ""
            for page in pdf.pages:
                text += (page.extract_text() or "") + "\n"

        bruto_kg = 0.0
        neto_kg = 0.0
        zemlja = ""

        m = _BRUTO_RE.search(text)
        if m:
            raw = m.group(1).replace(" ", "").replace(",", ".")
            try:
                bruto_kg = float(raw)
            except ValueError:
                pass

        m = _NETO_RE.search(text)
        if m:
            raw = m.group(1).replace(" ", "").replace(",", ".")
            try:
                neto_kg = float(raw)
            except ValueError:
                pass

        # Zemlja porijekla: traži samo jasno "MK" (OCR često kvari u "lrlK", "I\\4K", itd.)
        # Pekabesko je makedonski dobavljač; samo prihvatamo eksplicitan "MK".
        if re.search(r"poreklo\s*:?\s*MK\b", text, re.IGNORECASE):
            zemlja = "MK"

        logger.info(
            f"  📄 PDF ekstrakcija: bruto={bruto_kg:.2f}kg, "
            f"neto={neto_kg:.2f}kg, zemlja={zemlja}"
        )
        return bruto_kg, neto_kg, zemlja

    except Exception as e:
        logger.warning(f"  ⚠️ Greška pri PDF ekstrakciji: {e}")
        return 0.0, 0.0, ""


def parse_leburic_pekabesko_excel(filepath: str) -> ImportResult:
    """
    Parsira Leburic/Pekabesko Excel fakturu.

    Kolone (red 1 = header):
    A: Invoice, B: Item, C: Description, D: Unit, E: CustomID,
    F: Packets, G: Neto(KGR), H: Quantity In Unit, I: Price Per Unit,
    J: Total in EUR

    Returns:
        ImportResult sa stavkama, težinama i metapodacima
    """
    logger.info(f"Leburic/Pekabesko parsiranje: {filepath}")

    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.worksheets[0]

    # --- Mapiraj kolone iz headera ---
    header = {
        str(ws.cell(1, c).value or "").strip().lower(): c
        for c in range(1, ws.max_column + 1)
        if ws.cell(1, c).value
    }

    def col(name: str, *aliases) -> int:
        for n in (name, *aliases):
            if n.lower() in header:
                return header[n.lower()]
        return 0

    col_invoice   = col("invoice")
    col_item      = col("item")
    col_desc      = col("description")
    col_unit      = col("unit")
    col_custom    = col("customid")
    col_neto      = col("neto(kgr)")
    col_qty       = col("quantity in unit")
    col_price     = col("price per unit")
    col_total     = col("total in eur")

    # --- Pokušaj naći PDF i izvući bruto/zemlja ---
    bruto_kg = 0.0
    neto_kg_ukupno = 0.0
    zemlja_default = "MK"  # Pekabesko je uvijek iz Makedonije
    consumed_pdf: list[str] = []

    pdf_path = _find_pdf_for_excel(filepath)
    if pdf_path:
        bruto_pdf, neto_pdf, zemlja_pdf = _extract_from_pdf(pdf_path)
        if bruto_pdf > 0:
            bruto_kg = bruto_pdf
        if zemlja_pdf:
            zemlja_default = zemlja_pdf
        consumed_pdf = [pdf_path]  # PDF je iskorišten — agent ne treba da ga obrađuje ponovo

    # --- Čitaj stavke ---
    invoice_lines = []
    invoice_name = ""

    for row in range(2, ws.max_row + 1):
        item_val = ws.cell(row, col_item).value if col_item else None
        if not item_val:
            continue

        # Invoice broj (samo iz prvog reda)
        if not invoice_name and col_invoice:
            inv_raw = str(ws.cell(row, col_invoice).value or "").strip()
            if inv_raw:
                invoice_name = inv_raw

        product_code = str(item_val).strip()
        naziv = str(ws.cell(row, col_desc).value or "").strip().replace("_x000D_", "").strip() if col_desc else ""
        jm_raw = str(ws.cell(row, col_unit).value or "").strip() if col_unit else ""
        jm = _normalize_jm(jm_raw)

        tarifa = _read_tariff_code(ws.cell(row, col_custom).value if col_custom else None)

        neto_item = _parse_number(ws.cell(row, col_neto).value) if col_neto else 0.0
        kolicina = _parse_number(ws.cell(row, col_qty).value) if col_qty else 0.0
        cijena = _parse_number(ws.cell(row, col_price).value) if col_price else 0.0
        iznos = _parse_number(ws.cell(row, col_total).value) if col_total else 0.0

        neto_kg_ukupno += neto_item

        zemlja_iso = normalize_country_name(zemlja_default)

        line = InvoiceLine(
            line_no=len(invoice_lines) + 1,
            product_code=product_code,
            naziv_robe=naziv,
            tarifni_broj=tarifa,
            zemlja_porijekla=zemlja_iso,
            povlastica="",
            jm=jm,
            kolicina=kolicina,
            cijena_jed=cijena,
            iznos=iznos,
            valuta="EUR",
            bruto_kg=0.0,   # Nema po stavci; ukupni bruto je u ImportResult
            neto_kg=neto_item,
        )
        invoice_lines.append(line)

    wb.close()

    logger.info(
        f"  ✅ Parsed {len(invoice_lines)} stavki, "
        f"neto_sum={neto_kg_ukupno:.3f}kg, bruto={bruto_kg:.3f}kg"
    )

    return ImportResult(
        items=invoice_lines,
        bruto_kg=bruto_kg,
        neto_kg=neto_kg_ukupno,
        invoice_name=invoice_name,
        currency="EUR",
        import_type="leburic_pekabesko",
        exporter=Party(name="PEKABESKO"),
        consumed_paths=consumed_pdf,
    )
