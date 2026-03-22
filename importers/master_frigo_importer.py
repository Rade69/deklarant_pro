# importers/master_frigo_importer.py
"""
ASYCUDA Pro - Master Frigo Specialized Importer
Specijalizovani parser za Master Frigo fakture
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
import pdfplumber

from core.draft.draft import InvoiceLine
from importers.invoice_line_utils import parse_eu_number
from utils.country_normalizer import normalize_country_name

logger = logging.getLogger("asycuda_pro.import.master_frigo")


@dataclass
class ImportedLine:
    rbr: int
    code: str
    description: str
    unit: str
    qty: float
    price: float
    amount: float
    tariff: str = ""
    origin: str = ""
    preferential: str = ""
    serials: List[str] | None = None


_TABLE_HEADER_RE = re.compile(
    r"\bRbr\b.*\bSifra\b.*\bNaziv\b.*\bIznos\b", re.IGNORECASE
)
_INVOICE_NO_RE = re.compile(r"\bbroj:\s*([0-9]+)\b", re.IGNORECASE)
_DATE_RE = re.compile(r"\b(\d{1,2}/\d{1,2}/\d{4})\b")
_TOTAL_VALUE_RE = re.compile(
    r"\bVrednost\s*\((?P<cur>[A-Z]{3})\)\s*:\s*(?P<val>[\d\.,]+)", re.IGNORECASE
)
_GROSS_RE = re.compile(r"\bBruto težina:\s*(?P<val>[\d\.,]+)\s*KG\b", re.IGNORECASE)
_NET_RE = re.compile(r"\bNeto težina:\s*(?P<val>[\d\.,]+)\s*KG\b", re.IGNORECASE)
_INCOTERM_RE = re.compile(r"\bParitet isporuke:\s*(?P<term>[A-Z]{3})\b", re.IGNORECASE)

# unit + qty + price + amount
# Alternacija: ili 1 slovo (+ opcionalni razmak + 1-3 slova za split JM poput "k om"),
# ili direktno 2-3 slova. Ovo sprječava da 2-slovna oznaka modela (npr. "RX" u
# "GACC RX kom 1.00") bude greškom prepoznata kao prefiks JM.
_TAIL_NUMS_RE = re.compile(
    r"\s(?P<unit>[A-Za-z](?:\s+[A-Za-z]{1,3})?|[A-Za-z]{2,3})\s+"
    r"(?P<qty>[\d\.,]+)\s+(?P<price>[\d\.,]+)\s+(?P<amount>[\d\.,]+)\b"
)




def _read_tariff_code(value) -> str:
    """
    Čita tarifni broj iz Excel ćelije uz čuvanje vodećih nula.
    TARIC kodovi su 8 ili 10 cifara; 9/7 cifara znači izgubljena vodeća nula.
    """
    if value is None:
        return ""
    if isinstance(value, float):
        value = int(value)
    s = str(value).strip().replace(" ", "")
    if not s or s == "0":
        return ""
    if s.isdigit():
        n = len(s)
        if n == 9:
            s = s.zfill(10)
        elif n == 7:
            s = s.zfill(8)
    return s


def _extract_text_lines(pdf_path: str) -> List[str]:
    with pdfplumber.open(pdf_path) as pdf:
        lines: List[str] = []
        for page in pdf.pages:
            txt = page.extract_text() or ""
            lines.extend([ln.rstrip() for ln in txt.splitlines()])
    return lines


def _normalize_header_name(name: str) -> str:
    name = (name or "").strip().lower()
    name = (
        name.replace("š", "s")
        .replace("đ", "dj")
        .replace("č", "c")
        .replace("ć", "c")
        .replace("ž", "z")
    )
    name = re.sub(r"\s+", " ", name)
    return name


def _normalize_preferential(raw: str) -> str:
    """Povlastica/Preferencijal normalizacija"""
    s = (raw or "").strip()
    if not s:
        return ""
    u = s.upper()
    if u in {"DA", "YES", "Y", "1", "TRUE", "T"}:
        return "P"
    if u in {"NE", "NO", "N", "0", "FALSE", "F"}:
        return ""
    return s


def _read_mapping_xlsx(xlsx_path: str) -> Dict[str, Dict[str, str]]:
    """Čita mapping Excel: Šifra → tarifa/porijeklo/povlastica"""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    try:
        ws = wb[wb.sheetnames[0]]

        header_row = [str(ws.cell(1, c).value or "").strip() for c in range(1, 40)]
        idx_norm = {
            _normalize_header_name(name): i + 1 for i, name in enumerate(header_row) if name
        }

        def find_col(*candidates: str) -> Optional[int]:
            for cand in candidates:
                c = idx_norm.get(_normalize_header_name(cand))
                if c:
                    return c
            return None

        code_col = find_col("Šifra", "Sifra") or 2
        tariff_col = find_col("Tarifni br", "Tarifni broj") or 6
        origin_col = find_col("Zemlja porekla", "Zemlja porijekla", "Zemlja por") or 7
        pref_col = find_col("Preferencijal", "Povlastica", "Povlašćica") or 8

        out: Dict[str, Dict[str, str]] = {}
        for r in range(2, ws.max_row + 1):
            code = ws.cell(r, code_col).value
            if not code:
                continue

            code_s = str(code).strip().replace("\n", "").replace("\r", "")
            out[code_s] = {
                "tariff": _read_tariff_code(ws.cell(r, tariff_col).value),
                "origin": str(ws.cell(r, origin_col).value or "").strip(),
                "preferential": _normalize_preferential(
                    str(ws.cell(r, pref_col).value or "")
                ),
            }

        return out
    finally:
        # Ensure workbook is always closed
        wb.close()


def _find_code_and_desc(
    prefix_after_rbr: str, known_codes_sorted: List[str]
) -> Tuple[str, str]:
    """Pronalazi kod i opis (kod može imati razmake)"""
    s = prefix_after_rbr.strip()
    for code in known_codes_sorted:
        if s.startswith(code + " "):
            return code, s[len(code) :].strip()
        if s == code:
            return code, ""

    parts = s.split(None, 1)
    if not parts:
        return "", ""
    if len(parts) == 1:
        return parts[0], ""
    return parts[0], parts[1]


def parse_master_frigo_pdf(
    pdf_path: str, mapping: Optional[Dict[str, Dict[str, str]]] = None
) -> Tuple[Dict[str, Any], List[ImportedLine]]:
    """Parsira Master Frigo PDF fakturu"""
    lines = _extract_text_lines(pdf_path)
    
    # Ekstraktuj full tekst za detekciju izjave
    full_text = "\n".join(lines)

    header: Dict[str, Any] = {
        "source_file": Path(pdf_path).name,
        "has_origin_statement": False,
    }
    
    # DETEKTUJ SVE izjave o poreklu
    origin_statements = _detect_all_origin_statements(full_text)
    header["has_origin_statement"] = len(origin_statements) > 0
    header["origin_statements"] = origin_statements
    logger.info(f"  ✅ Detekcija izjave o poreklu: {header['has_origin_statement']} ({len(origin_statements)} izjava)")

    # Extract invoice number
    for ln in lines[:80]:
        m = _INVOICE_NO_RE.search(ln)
        if m:
            header["invoice_no"] = m.group(1)
            break

    # Extract date
    for ln in lines[:160]:
        if "Mesto i datum izdavanja" in ln:
            m = _DATE_RE.search(ln)
            if m:
                header["date"] = m.group(1)
                break
    if "date" not in header:
        for ln in lines[:160]:
            m = _DATE_RE.search(ln)
            if m:
                header["date"] = m.group(1)
                break

    # Extract currency and total - preferiramo EUR nad RSD
    # Faktura ima i RSD i EUR redove; EUR je prava valuta transakcije
    for ln in lines:
        m = _TOTAL_VALUE_RE.search(ln)
        if m:
            cur = m.group("cur").upper()
            val = parse_eu_number(m.group("val"))
            if "currency" not in header or cur == "EUR":
                header["currency"] = cur
                header["total_value"] = val
            if header.get("currency") == "EUR":
                break

    # Extract weights
    for ln in lines:
        m = _GROSS_RE.search(ln)
        if m:
            header["gross_kg"] = parse_eu_number(m.group("val"))
            break

    for ln in lines:
        m = _NET_RE.search(ln)
        if m:
            header["net_kg"] = parse_eu_number(m.group("val"))
            break

    # Extract incoterm
    for ln in lines:
        m = _INCOTERM_RE.search(ln)
        if m:
            header["incoterm"] = (m.group("term") or "").upper()
            break

    # Parse items
    start_idx: Optional[int] = None
    for i, ln in enumerate(lines):
        if _TABLE_HEADER_RE.search(ln):
            start_idx = i + 1
            break
    if start_idx is None:
        return header, []

    mapping = mapping or {}
    known_codes_sorted = sorted(mapping.keys(), key=len, reverse=True)

    items: List[ImportedLine] = []
    current: Optional[ImportedLine] = None
    collecting_serials = False
    skip_page_header = False  # True između "Licenca:" i sljedećeg "Rbr Sifra..." headera

    for ln in lines[start_idx:]:
        s = (ln or "").strip()
        if not s:
            continue

        # Preskoči ponavljanje zaglavlja stranice (Licenca: ... Rbr Sifra... Iznos)
        if s.lower().startswith("licenca:"):
            skip_page_header = True
            continue
        if skip_page_header:
            if _TABLE_HEADER_RE.search(s):
                skip_page_header = False
            continue

        if (
            s.upper().startswith("VREDNOST")
            or s.upper().startswith("UKUPNO")
            or "INSTRUCTION FOR" in s.upper()
        ):
            break

        if s.lower().startswith("serijski broj") or s.lower().startswith(
            "serijski brojevi"
        ):
            collecting_serials = True
            continue

        if collecting_serials and re.fullmatch(r"[0-9A-Z\-]{6,}", s.replace(" ", "")):
            if current:
                if current.serials is None:
                    current.serials = []
                current.serials.append(s.replace(" ", ""))
            continue

        if collecting_serials and re.match(r"^\d+\s", s):
            collecting_serials = False

        if re.match(r"^\d+\s", s):
            matches = list(_TAIL_NUMS_RE.finditer(s))
            if not matches:
                if current:
                    current.description = (current.description + " " + s).strip()
                continue

            m = matches[-1]
            unit = (m.group("unit") or "").replace(" ", "").lower()
            qty = parse_eu_number(m.group("qty"))
            price = parse_eu_number(m.group("price"))
            amount = parse_eu_number(m.group("amount"))
            tail = (s[m.end() :] or "").strip()

            prefix = (s[: m.start()] or "").strip()
            rbr_m = re.match(r"^(?P<rbr>\d+)\s+(?P<rest>.*)$", prefix)
            if not rbr_m:
                continue

            rbr = int(rbr_m.group("rbr"))
            rest = rbr_m.group("rest").strip()

            code, desc = _find_code_and_desc(rest, known_codes_sorted)
            if tail:
                desc = (desc + " " + tail).strip()

            map_row = mapping.get(code, {})
            current = ImportedLine(
                rbr=rbr,
                code=code,
                description=desc.strip(),
                unit=unit,
                qty=qty,
                price=price,
                amount=amount,
                tariff=map_row.get("tariff", "") or "",
                origin=map_row.get("origin", "") or "",
                preferential=_normalize_preferential(
                    map_row.get("preferential", "") or ""
                ),
                serials=[],
            )
            items.append(current)
        else:
            if current and not s.startswith(("56:", "57:", "59:")):
                # Provjeri da li red nastavlja nepotpunu šifru (npr. "GT-" + "GACC050.2/...")
                # PDF ponekad prelomi šifru u dvije linije; Excel ima kompletnu šifru
                code_completed = False
                if current.code.endswith("-") and known_codes_sorted:
                    # Probaj dva formata spajanja:
                    # 1. "GT- " + "GACC..." = "GT- GACC..." (sa razmakom)
                    # 2. "GT" + "GACC..." = "GTGACC..." (direktno spajanje bez crtice)
                    prefix_with_space = current.code + " "
                    prefix_no_dash = current.code[:-1]  # Ukloni krajnju crticu
                    for known_code in known_codes_sorted:
                        matched_prefix = None
                        if known_code.startswith(prefix_with_space):
                            matched_prefix = prefix_with_space
                        elif known_code.startswith(prefix_no_dash):
                            matched_prefix = prefix_no_dash
                        if matched_prefix:
                            suffix = known_code[len(matched_prefix):]
                            if s.startswith(suffix):
                                current.code = known_code
                                remainder = s[len(suffix):].strip()
                                if remainder:
                                    current.description = (current.description + " " + remainder).strip()
                                # Ažuriraj tariff/porijeklo za kompletnu šifru
                                full_map = mapping.get(known_code, {})
                                if full_map.get("tariff"):
                                    current.tariff = full_map["tariff"]
                                if full_map.get("origin"):
                                    current.origin = full_map["origin"]
                                if full_map.get("preferential"):
                                    current.preferential = _normalize_preferential(
                                        full_map["preferential"]
                                    )
                                code_completed = True
                                break
                if not code_completed:
                    current.description = (current.description + " " + s).strip()

    return header, items


def convert_to_invoice_lines(
    imported_lines: List[ImportedLine], currency: str = "EUR"
) -> List[InvoiceLine]:
    """Konvertuje ImportedLine → InvoiceLine"""
    invoice_lines = []

    for idx, item in enumerate(imported_lines):
        invoice_line = InvoiceLine(
            line_no=idx + 1,
            product_code=item.code,  # Za assembly matching po šifri
            naziv_robe=item.description,
            tarifni_broj=item.tariff,
            zemlja_porijekla=normalize_country_name(item.origin),  # Normalize to ISO code
            povlastica=item.preferential,
            jm=item.unit,
            kolicina=item.qty,
            cijena_jed=item.price,
            iznos=item.amount,
            valuta=currency,
            bruto_kg=0.0,  # Nema pojedinačnih težina po stavci
            neto_kg=0.0,
        )
        invoice_lines.append(invoice_line)

    return invoice_lines


def import_master_frigo(
    pdf_path: str, mapping_xlsx_path: Optional[str] = None
) -> List[InvoiceLine]:
    """
    Import Master Frigo fakture.

    Args:
        pdf_path: Putanja do PDF fakture
        mapping_xlsx_path: Opciono - Excel mapping za tarife/porijeklo

    Returns:
        List[InvoiceLine]
    """
    logger.info(f"Importing Master Frigo PDF: {pdf_path}")

    # Load mapping if provided
    mapping = _read_mapping_xlsx(mapping_xlsx_path) if mapping_xlsx_path else {}

    # Parse PDF
    header, imported_items = parse_master_frigo_pdf(pdf_path, mapping=mapping)

    # Convert to InvoiceLine
    currency = header.get("currency", "EUR")
    invoice_lines = convert_to_invoice_lines(imported_items, currency=currency)

    logger.info(f"Imported {len(invoice_lines)} items from Master Frigo PDF")

    return invoice_lines


def _detect_all_origin_statements(text: str) -> list:
    """
    Detektuj SVE izjave o preferencijalnom poreklu u tekstu.

    Koristi OriginStatementDetector servis.

    Args:
        text: Tekst PDF fakture

    Returns:
        Lista OriginStatementMatch objekata
    """
    try:
        from services.origin_statement_detector import OriginStatementDetector

        detector = OriginStatementDetector()
        statements = detector.detect_all_in_text(text)

        for stmt in statements:
            logger.info(f"  ✅ Nađena izjava: {stmt.jezik} / {stmt.tip_izjave} / origin={stmt.origin_country}")

        return statements

    except Exception as e:
        logger.warning(f"  ⚠️  Greška tokom detekcije izjava: {e}")
        return []
