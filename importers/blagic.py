# services/invoice_importers/blagic.py
from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
import re

import openpyxl
import pdfplumber


_INVOICE_ID_RE = re.compile(r"(\d+VP-\d{4})", re.IGNORECASE)


@dataclass
class ImportedLine:
    num: int
    code: str
    description: str
    unit: str
    qty: float
    price: float
    amount: float
    tariff: str = ""
    origin: str = ""
    weight_total: float = 0.0
    weight_piece: float = 0.0
    packing_matched: bool = False


def extract_invoice_id_from_filename(path: str) -> Optional[str]:
    m = _INVOICE_ID_RE.search(Path(path).name)
    return m.group(1).upper() if m else None


def _normalize_pdf_number(s: str) -> float:
    """
    PDF u ovom formatu koristi decimalni zarez.
    Primjer: '1,20' -> 1.2
    """
    s = (s or "").strip()
    s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def _parse_decimal_any(s: str) -> Decimal:
    s = (s or "").strip()
    s = s.replace(" ", "")
    # podrži i '2.459,08' i '2459,08' i '2459.08'
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    else:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return Decimal("0")


def _format_decimal(d: Decimal, places: int = 2) -> str:
    q = Decimal("1." + ("0" * places))
    return str(d.quantize(q))


def _looks_like_service_line(desc: str) -> bool:
    """
    U praksi: troškovi špedicije/transporta/usluga nisu roba.
    Ove stavke izbacujemo iz NAIMENOVANJA, ali ostaju u fakturi (ukupno).
    """
    d = (desc or "").upper()
    keywords = [
        "TROSK",
        "TROŠK",
        "ŠPED",
        "SPED",
        "USLUG",
        "TRANSPORT",
        "PREVOZ",
        "CARIN",
        "BANK",
        "PROVIZ",
    ]
    return any(k in d for k in keywords)


def parse_blagic_invoice_pdf(
    pdf_path: str,
) -> Tuple[Dict[str, Any], List[ImportedLine]]:
    """
    Parsira BLAGIĆ/Loren PDF fakturu:
      - Header: invoice_no, invoice_date (dd.mm.yyyy), currency (EUR), supplier, customer
      - Stavke: num, code, description, unit, qty, price, amount
    """
    with pdfplumber.open(pdf_path) as pdf:
        lines: List[str] = []
        for page in pdf.pages:
            text = page.extract_text() or ""
            lines.extend(text.splitlines())

    header: Dict[str, Any] = {"currency": "EUR"}

    # supplier: prva ne-prazna linija
    for ln in lines:
        if ln.strip():
            header["supplier"] = ln.strip()
            break

    # invoice no / date / customer
    for ln in lines:
        m = re.search(r"Invoice:\s*([A-Z0-9\-]+)", ln)
        if m:
            header["invoice_no"] = m.group(1).strip()

        m = re.search(r"Invoice date:\s*([0-9]{1,2}\.[a-zA-Z]{3}\.[0-9]{4}\.)", ln)
        if m:
            header["invoice_date_raw"] = m.group(1).strip()

        if "Address:" in ln and "Beograd" in ln:
            m2 = re.search(r"\b11000\s+Beograd\s+(.*)$", ln)
            if m2 and m2.group(1).strip():
                header["customer"] = m2.group(1).strip()

    # invoice_date: 18.nov.2025. -> 18.11.2025
    if header.get("invoice_date_raw"):
        raw = header["invoice_date_raw"].strip(".")
        parts = raw.split(".")
        if len(parts) >= 3:
            day = int(parts[0])
            mon = parts[1].lower()[:3]
            year = int(parts[2])
            months = {
                "jan": 1,
                "feb": 2,
                "mar": 3,
                "apr": 4,
                "may": 5,
                "jun": 6,
                "jul": 7,
                "aug": 8,
                "sep": 9,
                "oct": 10,
                "okt": 10,
                "nov": 11,
                "dec": 12,
            }
            header["invoice_date"] = f"{day:02d}.{months.get(mon, 1):02d}.{year:04d}"

    started = False
    current: Optional[ImportedLine] = None
    items: List[ImportedLine] = []

    def can_parse_tail(tokens: List[str]) -> bool:
        """
        Očekujemo da su zadnja 3 tokena broj sa zarezom: QTY, Price, Amount.
        """
        if len(tokens) < 4:
            return False
        qty, price, amount = tokens[-3], tokens[-2], tokens[-1]
        for x in (qty, price, amount):
            if not re.match(r"^\d+,\d{2}$", x):
                return False
        return True

    for ln in lines:
        if "Num." in ln and "Code" in ln and "Amount" in ln:
            started = True
            continue
        if not started:
            continue

        s = ln.strip()
        if not s:
            continue

        # stop na footer/summary
        if (
            "STRANA" in s.upper()
            or s.upper().startswith("AMOUNT:")
            or "AMOUNT:" in s.upper()
        ):
            break
        if "TOTAL" in s.upper():
            break

        m = re.match(r"^\s*(\d+)\s+([A-Z0-9]+)\s+(.*)$", ln)
        if m:
            num = int(m.group(1))
            code = m.group(2).strip()
            rest = m.group(3).strip()
            tokens = rest.split()

            if can_parse_tail(tokens):
                unit = tokens[-4]
                qty = _normalize_pdf_number(tokens[-3])
                price = _normalize_pdf_number(tokens[-2])
                amount = _normalize_pdf_number(tokens[-1])
                desc = " ".join(tokens[:-4]).strip()

                current = ImportedLine(
                    num=num,
                    code=code,
                    description=desc,
                    unit=unit,
                    qty=qty,
                    price=price,
                    amount=amount,
                )
                items.append(current)
                continue

            # ako rep ne može da se parsira: nastavak opisa
            if current:
                current.description += " " + s
            continue

        # continuation line
        if current:
            current.description += " " + s

    return header, items


def parse_blagic_packing_xlsx(xlsx_path: str) -> Dict[str, Dict[str, Any]]:
    """
    Excel packing lista sa kolonama:
      RB | Kod | Artikal | JM | Kolicina | Težina/kom | Težina ukupno | Tarifni broj | Poreklo

    Vraća dict: code -> data
    """
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    out: Dict[str, Dict[str, Any]] = {}

    for r in range(2, ws.max_row + 1):
        rb = ws.cell(r, 1).value
        code = ws.cell(r, 2).value
        if rb is None or not code:
            continue

        code_s = str(code).strip()

        out[code_s] = {
            "rb": int(rb) if isinstance(rb, (int, float)) else rb,
            "code": code_s,
            "article": ws.cell(r, 3).value or "",
            "unit": ws.cell(r, 4).value or "",
            "qty": float(ws.cell(r, 5).value or 0),
            "weight_piece": float(ws.cell(r, 6).value or 0),
            "weight_total": float(ws.cell(r, 7).value or 0),
            "tariff": str(ws.cell(r, 8).value or "").strip(),
            "origin": str(ws.cell(r, 9).value or "").strip(),
        }

    return out


def merge_invoice_with_packing(
    invoice_items: List[ImportedLine], packing_by_code: Dict[str, Dict[str, Any]]
) -> List[ImportedLine]:
    merged: List[ImportedLine] = []
    for it in invoice_items:
        pack = packing_by_code.get(it.code)
        if pack:
            it.tariff = pack.get("tariff", "") or ""
            it.origin = pack.get("origin", "") or ""
            it.weight_total = float(pack.get("weight_total") or 0.0)
            it.weight_piece = float(pack.get("weight_piece") or 0.0)
            it.packing_matched = True
        merged.append(it)
    return merged


def build_faktura_tab_data(
    header: Dict[str, Any], merged_lines: List[ImportedLine]
) -> Dict[str, Any]:
    total = sum(Decimal(str(x.amount or 0)) for x in merged_lines)
    return {
        "broj_fakture": header.get("invoice_no", "") or "",
        "datum_fakture": header.get("invoice_date", "") or "",
        "valuta": header.get("currency", "EUR") or "EUR",
        "ukupna_vrijednost": _format_decimal(total, 2),
        "dobavljac": header.get("supplier", "") or "",
        "kupac": header.get("customer", "") or "",
    }


def _join_full_descriptions(desc_list: List[str]) -> str:
    """
    Puni opis za rub.31:
      - unique (po redu pojavljivanja)
      - svaki opis u novi red (čita se lako u rub.31)
    """
    seen = set()
    uniq: List[str] = []
    for d in desc_list:
        d = (d or "").strip()
        if not d:
            continue
        if d in seen:
            continue
        seen.add(d)
        uniq.append(d)
    return "\n".join(uniq)


def build_naimenovanja_rows(merged_lines: List[ImportedLine]) -> List[List[Any]]:
    """
    Treeview redovi za NAIMENOVANJA:
      [rb, tarifni, opis(31 FULL), kolicina, vrijednost, bruto, neto, porijeklo, povlastica]

    Grupisanje: (tariff, origin)
    Service stavke se izbacuju.
    """
    from collections import defaultdict

    groups = defaultdict(
        lambda: {"qty": 0.0, "value": 0.0, "net": 0.0, "gross": 0.0, "desc": []}
    )

    for it in merged_lines:
        if _looks_like_service_line(it.description):
            continue

        key = (it.tariff or "", it.origin or "")
        g = groups[key]
        g["qty"] += float(it.qty or 0)
        g["value"] += float(it.amount or 0)
        g["net"] += float(it.weight_total or 0)
        g["gross"] += float(it.weight_total or 0)
        g["desc"].append(it.description or "")

    rows: List[List[Any]] = []
    for idx, ((tariff, origin), g) in enumerate(
        sorted(groups.items(), key=lambda x: (x[0][0], x[0][1])), start=1
    ):
        full_desc_31 = _join_full_descriptions(g["desc"])  # NEMA truncation-a
        rows.append(
            [
                idx,
                tariff,
                full_desc_31,
                round(g["qty"], 3),
                round(g["value"], 2),
                round(g["gross"], 3),
                round(g["net"], 3),
                origin,
                "",
            ]
        )
    return rows


def import_blagic_pair(invoice_pdf_path: str, packing_xlsx_path: str) -> Dict[str, Any]:
    header, invoice_items = parse_blagic_invoice_pdf(invoice_pdf_path)
    packing = parse_blagic_packing_xlsx(packing_xlsx_path)
    merged = merge_invoice_with_packing(invoice_items, packing)

    faktura_data = build_faktura_tab_data(header, merged)
    naimenovanja_data = {
        "items": build_naimenovanja_rows(merged),
        "editor": {},
    }

    raw_lines = [it.__dict__ for it in merged]

    missing_tariffs = sorted(
        {
            it.code
            for it in merged
            if not _looks_like_service_line(it.description)
            and not (it.tariff or "").strip()
        }
    )

    return {
        "header": header,
        "faktura": faktura_data,
        "naimenovanja": naimenovanja_data,
        "raw_lines": raw_lines,
        "missing_tariffs": missing_tariffs,
    }


def combine_imports(import_results: List[Dict[str, Any]]) -> Dict[str, Any]:
    """
    Kombinuje više uvoza u jedan projekat:
      - faktura: suma ukupne vrijednosti, broj_fakture = join
      - naimenovanja: merge grupa po (tariff, origin)
      - opis: puni opis (više linija) se čuva u row[2]
    """
    if not import_results:
        return {
            "faktura": {},
            "naimenovanja": {"items": [], "editor": {}},
            "missing_tariffs": [],
        }

    broj_list: List[str] = []
    datum = import_results[0].get("faktura", {}).get("datum_fakture", "")
    valuta = import_results[0].get("faktura", {}).get("valuta", "EUR")
    dob = import_results[0].get("faktura", {}).get("dobavljac", "")
    kup = import_results[0].get("faktura", {}).get("kupac", "")

    total = Decimal("0")

    for res in import_results:
        f = res.get("faktura", {})
        if f.get("broj_fakture"):
            broj_list.append(f["broj_fakture"])
        total += _parse_decimal_any(str(f.get("ukupna_vrijednost", "0")))

    faktura = {
        "broj_fakture": ", ".join(broj_list),
        "datum_fakture": datum,
        "valuta": valuta,
        "ukupna_vrijednost": _format_decimal(total, 2),
        "dobavljac": dob,
        "kupac": kup,
    }

    from collections import defaultdict

    agg = defaultdict(
        lambda: {"qty": 0.0, "value": 0.0, "gross": 0.0, "net": 0.0, "desc_list": []}
    )

    for res in import_results:
        rows = res.get("naimenovanja", {}).get("items", [])
        for r in rows:
            if not isinstance(r, list) or len(r) < 9:
                continue
            tariff = str(r[1] or "")
            full_desc = str(r[2] or "")
            qty = float(r[3] or 0)
            val = float(r[4] or 0)
            gross = float(r[5] or 0)
            net = float(r[6] or 0)
            origin = str(r[7] or "")
            key = (tariff, origin)
            agg[key]["qty"] += qty
            agg[key]["value"] += val
            agg[key]["gross"] += gross
            agg[key]["net"] += net
            if full_desc.strip():
                # full_desc može već biti višelinijski; sačuvaj kao “blok”
                agg[key]["desc_list"].append(full_desc.strip())

    combined_rows = []
    for idx, ((tariff, origin), g) in enumerate(
        sorted(agg.items(), key=lambda x: (x[0][0], x[0][1])), start=1
    ):
        # spajanje blokova opisa: svaki blok odvoji praznom linijom
        # (rub.31 čitljiv, bez truncation-a)
        full31 = "\n\n".join([d for d in g["desc_list"] if d])
        combined_rows.append(
            [
                idx,
                tariff,
                full31,
                round(g["qty"], 3),
                round(g["value"], 2),
                round(g["gross"], 3),
                round(g["net"], 3),
                origin,
                "",
            ]
        )

    missing = sorted(
        {code for res in import_results for code in res.get("missing_tariffs", [])}
    )

    return {
        "faktura": faktura,
        "naimenovanja": {"items": combined_rows, "editor": {}},
        "missing_tariffs": missing,
    }
